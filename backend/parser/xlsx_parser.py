import io
import openpyxl
from typing import List, Tuple
from app.models import Block, MetadataItem

def parse_xlsx(content: bytes) -> Tuple[List[Block], List[MetadataItem], List[str], int]:
    blocks: List[Block] = []
    metadata: List[MetadataItem] = []
    warnings: List[str] = []
    hidden_count = 0
    order = 0

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # Core properties
    props = wb.properties
    if props.creator:
        metadata.append(MetadataItem(key="Creator", value=props.creator, source="workbook_properties"))
        hidden_count += 1
    if props.lastModifiedBy:
        metadata.append(MetadataItem(key="Last Modified By", value=props.lastModifiedBy, source="workbook_properties"))
        hidden_count += 1

    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        is_sheet_hidden = sheet.sheet_state != "visible"
        if is_sheet_hidden:
            hidden_count += 1
            warnings.append(f"Sheet '{sheetname}' is hidden ({sheet.sheet_state}).")

        # Iterate rows
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
            if not row_vals:
                continue
            
            row_text = f"[{sheetname} Row {r_idx}] " + " | ".join(row_vals)
            blocks.append(Block(
                id=f"xlsx_{sheetname}_{r_idx}",
                type="cell",
                text=row_text,
                visibility="hidden" if is_sheet_hidden else "visible",
                locator={"sheet": sheetname, "row": r_idx, "sheet_state": sheet.sheet_state},
                order=order
            ))
            order += 1

    return blocks, metadata, warnings, hidden_count
