import io
import docx
from typing import List, Tuple
from app.models import Block, MetadataItem

def parse_docx(content: bytes) -> Tuple[List[Block], List[MetadataItem], List[str], int]:
    blocks: List[Block] = []
    metadata: List[MetadataItem] = []
    warnings: List[str] = []
    hidden_count = 0
    order = 0

    doc = docx.Document(io.BytesIO(content))

    # Core metadata
    core_props = doc.core_properties
    if core_props.author:
        metadata.append(MetadataItem(key="Author", value=core_props.author, source="core_properties"))
        hidden_count += 1
    if core_props.last_modified_by:
        metadata.append(MetadataItem(key="Last Modified By", value=core_props.last_modified_by, source="core_properties"))
        hidden_count += 1
    if core_props.title:
        metadata.append(MetadataItem(key="Title", value=core_props.title, source="core_properties"))

    # Paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        
        # Check style for headings
        b_type = "heading" if p.style.name.startswith("Heading") else "paragraph"
        blocks.append(Block(
            id=f"docx_p_{p_idx}",
            type=b_type,
            text=text,
            visibility="visible",
            locator={"paragraph": p_idx, "style": p.style.name},
            order=order
        ))
        order += 1

    # Tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip() for cell in row.cells]
            combined_row = " | ".join([t for t in row_texts if t])
            if combined_row:
                blocks.append(Block(
                    id=f"docx_tbl_{t_idx}_r_{r_idx}",
                    type="table",
                    text=combined_row,
                    visibility="visible",
                    locator={"table": t_idx, "row": r_idx},
                    order=order
                ))
                order += 1

    if hidden_count > 0:
        warnings.append(f"Document contains {hidden_count} metadata entries (Author/Modified By) that will be screened.")

    return blocks, metadata, warnings, hidden_count
